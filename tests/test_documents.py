import io
import zipfile
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from rgce import core, documents
from rgce.app import create_app
from rgce.store import Store


def test_text_receipt_dedup_and_source_citations(tmp_path):
    store = Store(tmp_path)
    r = documents.intake(store, "policy.md", b"# Sample policy\nUse private networks.", "policies", "organization", "Sample team")
    assert r["body"]["status"] == "unreviewed"
    assert r["body"]["origin"] == "source_document"
    assert documents.intake(store, "policy.md", b"# Sample policy\nUse private networks.", "policies", "organization", "Sample team")["id"] == r["id"]
    assert (tmp_path / "knowledge/documents" / r["id"] / "content.md").exists()
    assert documents.search(store, "private", "policies")["results"][0]["id"] == r["id"]
    a = core.demo(store)["before"]
    draft = core.draft_knowledge(store, a["id"], "The policy requires private networks.", "Reviewer", [r["id"]])
    approved = core.approve_knowledge(store, draft["id"], draft["id"], "Reviewer")
    assert approved["body"]["source_ids"] == [r["id"]]
    assert store.get(r["id"])["body"]["status"] == "unreviewed"


@pytest.mark.parametrize("name,data", [("../policy.md", b"text"), ("https://host/a.md", b"text"), ("bundle.zip", b"PK"), ("fake.pdf", b"text"), ("macro.docm", b"PK"), ("big.md", b"a" * 5_000_001)], ids=["traversal", "url", "zip", "fake-pdf", "macro", "oversized"])
def test_reject_unbounded_or_unsupported_inputs(name, data):
    with pytest.raises(ValueError): documents.validate_file(name, data)


def test_macro_in_renamed_docx_rejected():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w") as z:
        z.writestr("word/document.xml", "<document/>")
        z.writestr("word/vbaProject.bin", "macro")
    with pytest.raises(ValueError, match="Macro"): documents.validate_file("renamed.docx", data.getvalue())


def test_real_markitdown_xlsx(tmp_path):
    from openpyxl import Workbook
    b = io.BytesIO()
    book = Workbook()
    book.active.append(["Service", "Owner"])
    book.active.append(["Sample API", "Sample team"])
    book.save(b)
    r = documents.intake(Store(tmp_path), "ownership.xlsx", b.getvalue(), "ownership")
    assert "Sample API" in r["body"]["markdown"]
    assert r["body"]["converter"].startswith("markitdown/")


@pytest.mark.parametrize("suffix", ["pdf", "docx"])
def test_real_markitdown_fixture(tmp_path, suffix):
    data = (Path(__file__).parent / "fixtures/knowledge" / ("sample." + suffix)).read_bytes()
    r = documents.intake(Store(tmp_path), "sample." + suffix, data)
    assert "private networks only" in r["body"]["markdown"]


def test_real_markitdown_pptx_and_csv(tmp_path):
    from pptx import Presentation
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "Sample architecture"
    slide.placeholders[1].text = "Private networks only"
    b = io.BytesIO()
    p.save(b)
    store = Store(tmp_path)
    r = documents.intake(store, "architecture.pptx", b.getvalue(), "architecture")
    assert "Private networks only" in r["body"]["markdown"]
    csv = documents.intake(store, "owners.csv", b"Service,Owner\nSample,Platform\n", "ownership")
    assert "Platform" in csv["body"]["markdown"]


def test_file_upload_needs_local_session(tmp_path):
    with TestClient(create_app(tmp_path)) as c:
        url = "/api/documents?name=runbook.md&topic=runbooks&scope=application:sample"
        assert c.post(url, content=b"# Sample").status_code == 403
        headers = {"X-RGCE-Token": c.get("/api/session").json()["token"]}
        response = c.post(url, content=b"# Sample runbook", headers=headers)
        assert response.status_code == 200
        assert c.get("/api/documents?query=runbook").json()["results"][0]["body"]["scope"] == "application:sample"
        assert not c.get("/api/state").json()["knowledge"]


def test_mcp_document_excerpt_is_bounded(tmp_path):
    import asyncio
    import json
    from mcp import Client
    from rgce.mcp_server import create_server
    store = Store(tmp_path)
    source = documents.intake(store, "long.md", b"sample context\n" * 2000)
    async def run():
        async with Client(create_server(store)) as client:
            result = await client.call_tool("get_record", {"record_id": source["id"], "limit": 100})
            body = result.structured_content or json.loads(result.content[0].text)
            assert len(body["excerpt"]) == 100 and body["next_offset"] == 100
            assert not body["complete_document"]
            assert body["metadata"]["status"] == "unreviewed"
    asyncio.run(run())
